"""Security boundary and document operations for the local MCP server."""
from __future__ import annotations

import hashlib
import hmac
import importlib.util
import io
import json
import os
import re
import secrets
import sqlite3
import stat
import subprocess
import sys
import threading
import time
import zipfile
from contextlib import contextmanager
from importlib.metadata import version as metadata_version, PackageNotFoundError
from dataclasses import dataclass
from pathlib import Path, PurePath
from typing import Protocol


class GatewayError(Exception):
    """Public, non-confidential operation failure."""


def _fail() -> GatewayError:
    return GatewayError("Operation denied or document unavailable")


@dataclass(frozen=True)
class Config:
    roots: tuple[Path, ...]
    workspace: Path
    language: str
    entities: tuple[str, ...]
    threshold: float
    ttl: int
    max_source: int
    max_output: int
    scope: str
    spacy_model: str | None = None
    policy_version: str = "1"

    @classmethod
    def load(cls, filename: str) -> "Config":
        raw = json.loads(Path(filename).read_text(encoding="utf-8"))
        required = {"source_roots", "workspace", "language", "entities", "score_threshold",
                    "cache_ttl_seconds", "max_source_bytes", "max_output_chars", "access_scope"}
        if not required.issubset(raw) or set(raw) - required - {"spacy_model", "policy_version"}:
            raise ValueError("Invalid configuration keys")
        roots = tuple(Path(p).absolute() for p in raw["source_roots"])
        workspace = Path(raw["workspace"]).absolute()
        if not roots or any(not p.is_absolute() or not p.is_dir() or p.is_symlink() for p in roots):
            raise ValueError("Invalid source roots")
        if not Path(raw["workspace"]).is_absolute() or not all(Path(p).is_absolute() for p in raw["source_roots"]):
            raise ValueError("Absolute paths required")
        resolved_roots = tuple(p.resolve(strict=True) for p in roots)
        wr = workspace.resolve()
        if any(wr == r or wr.is_relative_to(r) or r.is_relative_to(wr) for r in resolved_roots):
            raise ValueError("Workspace and sources must be disjoint")
        if workspace.exists() and (workspace.is_symlink() or not workspace.is_dir()):
            raise ValueError("Invalid workspace")
        language = raw["language"]
        entities = tuple(raw["entities"])
        scope = raw["access_scope"]
        spacy_model = raw.get("spacy_model")
        policy_version = raw.get("policy_version", "1")
        if not isinstance(language, str) or not re.fullmatch(r"[a-z]{2}", language):
            raise ValueError("Invalid language")
        if not entities or any(not isinstance(e, str) or not re.fullmatch(r"[A-Z_]+", e) for e in entities):
            raise ValueError("Invalid entities")
        if not isinstance(scope, str) or not re.fullmatch(r"[A-Za-z0-9_-]{1,64}", scope):
            raise ValueError("Invalid access scope")
        if spacy_model is not None and (not isinstance(spacy_model, str) or not re.fullmatch(r"[A-Za-z0-9_]+", spacy_model)):
            raise ValueError("Invalid spaCy model")
        if not spacy_model:
            raise ValueError("Explicit local spaCy model required")
        if not isinstance(policy_version, str) or not re.fullmatch(r"[A-Za-z0-9_.-]{1,64}", policy_version):
            raise ValueError("Invalid policy version")
        threshold = float(raw["score_threshold"])
        ttl = int(raw["cache_ttl_seconds"])
        max_source = int(raw["max_source_bytes"])
        max_output = int(raw["max_output_chars"])
        if not 0 <= threshold <= 1 or ttl < 1 or max_source < 1 or max_output < 1:
            raise ValueError("Invalid limits")
        return cls(resolved_roots, workspace, language, entities, threshold, ttl, max_source, max_output, scope, spacy_model, policy_version)

    def policy_digest(self) -> str:
        data = [self.language, self.entities, self.threshold, self.scope, self.spacy_model,
                [str(r) for r in self.roots], "text-v1", "pseudonym-v1", self.policy_version]
        return hashlib.sha256(json.dumps(data, sort_keys=True).encode()).hexdigest()


class Detector(Protocol):
    identity: str
    def detect(self, text: str, language: str, entities: tuple[str, ...], threshold: float) -> list[tuple[int, int, str]]: ...


class PresidioDetector:
    def __init__(self, language: str, spacy_model: str | None = None):
        try:
            if not spacy_model or importlib.util.find_spec(spacy_model) is None:
                raise RuntimeError("Local spaCy model unavailable")
            import spacy
            spacy.load(spacy_model)  # Local-only preflight before Presidio can attempt auto-download.
            from presidio_analyzer import AnalyzerEngine, Pattern, PatternRecognizer
            from presidio_analyzer.nlp_engine import NlpEngineProvider
            provider = NlpEngineProvider(nlp_configuration={"nlp_engine_name": "spacy", "models": [
                {"lang_code": language, "model_name": spacy_model}]})
            self.engine = AnalyzerEngine(nlp_engine=provider.create_engine(), supported_languages=[language])
            # Force model loading at startup. No silent regex-only fallback.
            self.engine.nlp_engine.load()
            if language == "pl":
                patterns = {"PESEL": r"\b\d{11}\b", "NIP": r"\b\d{10}\b", "REGON": r"\b(?:\d{9}|\d{14})\b"}
                for entity, regex in patterns.items():
                    self.engine.registry.add_recognizer(PatternRecognizer(supported_entity=entity,
                        patterns=[Pattern(name=entity, regex=regex, score=0.5)], supported_language="pl"))
            try:
                package_version = metadata_version("presidio-analyzer")
            except PackageNotFoundError:
                package_version = "unknown"
            self.identity = "presidio/" + package_version + "/" + language + "/" + spacy_model + "/" + metadata_version("spacy") + "/" + metadata_version(spacy_model.replace("_", "-"))
        except Exception as exc:
            raise RuntimeError("Presidio NLP model unavailable") from exc

    def detect(self, text: str, language: str, entities: tuple[str, ...], threshold: float) -> list[tuple[int, int, str]]:
        try:
            matches = self.engine.analyze(text=text, language=language, entities=list(entities), score_threshold=threshold)
            return [(m.start, m.end, m.entity_type) for m in matches]
        except Exception as exc:
            raise GatewayError("PII detection unavailable") from exc


def _extract(data: bytes, suffix: str) -> str:
    try:
        if suffix == ".txt":
            return data.decode("utf-8-sig", errors="strict")
        if suffix == ".docx":
            from docx import Document
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                _check_archive(archive)
                if any(n.startswith(("word/media/", "word/embeddings/")) for n in archive.namelist()):
                    raise _fail()
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                parts.extend(cell.text for row in table.rows for cell in row.cells)
            for section in doc.sections:
                parts.extend(p.text for p in section.header.paragraphs)
                parts.extend(p.text for p in section.footer.paragraphs)
            return "\n".join(parts)
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data), strict=True)
            if reader.is_encrypted:
                raise _fail()
            if len(reader.pages) > 1000:
                raise _fail()
            pages = []
            for page in reader.pages:
                if page.images:
                    raise _fail()
                value = page.extract_text()
                if not value or not value.strip():
                    raise _fail()
                pages.append(value)
            return "\n\n".join(pages)
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                _check_archive(archive)
                if any(n.startswith(("xl/media/", "xl/externalLinks/", "xl/embeddings/")) for n in archive.namelist()):
                    raise _fail()
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            lines = []
            cell_count = 0
            for sheet in book.worksheets:
                lines.append("[sheet]")
                for row in sheet.iter_rows():
                    cells = []
                    for cell in row:
                        cell_count += 1
                        if cell_count > 100000:
                            raise _fail()
                        if cell.data_type == "f":
                            raise _fail()
                        cells.append("" if cell.value is None else str(cell.value))
                    lines.append("\t".join(cells))
            book.close()
            return "\n".join(lines)
    except GatewayError:
        raise
    except Exception as exc:
        raise _fail() from exc
    raise _fail()


def _check_archive(archive: zipfile.ZipFile) -> None:
    members = archive.infolist()
    if len(members) > 1000 or sum(m.file_size for m in members) > 40_000_000:
        raise _fail()


_PARSER_SLOT = threading.BoundedSemaphore(1)


def _extract_bounded(data: bytes, suffix: str, max_output: int, max_source: int) -> str:
    """Parse untrusted binary formats in a time- and memory-limited process."""
    if not _PARSER_SLOT.acquire(timeout=20):
        raise _fail()
    try:
        result = subprocess.run(
            [sys.executable, "-m", "privacy_gateway.extractor", suffix, str(max_output), str(max_source)],
            input=data, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
            timeout=20, check=False,
        )
        if result.returncode != 0 or len(result.stdout) > max_output * 4:
            raise _fail()
        return result.stdout.decode("utf-8", errors="strict")
    except GatewayError:
        raise
    except Exception as exc:
        raise _fail() from exc
    finally:
        _PARSER_SLOT.release()


class Gateway:
    def __init__(self, config: Config, key: bytes, detector: Detector | None = None):
        from cryptography.fernet import Fernet
        if len(key) < 32:
            raise ValueError("Gateway key must contain at least 32 random bytes")
        self.config = config
        self.key = key
        self.fernet = Fernet(__import__("base64").urlsafe_b64encode(hashlib.sha256(key + b"encryption").digest()))
        self.detector = detector if detector is not None else PresidioDetector(config.language, config.spacy_model)
        os.umask(0o077)
        config.workspace.mkdir(mode=0o700, parents=True, exist_ok=True)
        if config.workspace.is_symlink():
            raise ValueError("Workspace symlink")
        os.chmod(config.workspace, 0o700)
        self.db_path = config.workspace / "gateway.sqlite3"
        with self._db() as db:
            db.execute("CREATE TABLE IF NOT EXISTS documents (id TEXT PRIMARY KEY, source BLOB NOT NULL, kind TEXT NOT NULL, version INTEGER NOT NULL)")
            db.execute("CREATE TABLE IF NOT EXISTS cache (cache_key TEXT PRIMARY KEY, payload BLOB NOT NULL, expires INTEGER NOT NULL)")

    @contextmanager
    def _db(self):
        db = sqlite3.connect(self.db_path, timeout=10)
        db.execute("PRAGMA busy_timeout=10000")
        try:
            yield db
            db.commit()
        except Exception:
            db.rollback()
            raise
        finally:
            db.close()

    def _seal(self, value: bytes, context: str) -> bytes:
        return self.fernet.encrypt(context.encode() + b"\0" + value)

    def _open_seal(self, value: bytes, context: str) -> bytes:
        try:
            decoded = self.fernet.decrypt(value)
            prefix, payload = decoded.split(b"\0", 1)
            if not hmac.compare_digest(prefix, context.encode()):
                raise _fail()
            return payload
        except Exception as exc:
            raise _fail() from exc

    def _source_bytes(self, root_index: int, relative: str) -> bytes:
        if root_index < 0 or root_index >= len(self.config.roots):
            raise _fail()
        parts = PurePath(relative).parts
        if not parts or any(p in (".", "..", "") for p in parts) or PurePath(relative).is_absolute():
            raise _fail()
        descriptors = []
        try:
            current = os.open("/", os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)
            descriptors.append(current)
            for part in self.config.roots[root_index].parts[1:]:
                current = os.open(part, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW, dir_fd=current)
                descriptors.append(current)
            for part in parts[:-1]:
                current = os.open(part, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW, dir_fd=current)
                descriptors.append(current)
            fd = os.open(parts[-1], os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK, dir_fd=current)
            descriptors.append(fd)
            metadata = os.fstat(fd)
            if not stat.S_ISREG(metadata.st_mode) or metadata.st_size > self.config.max_source:
                raise _fail()
            with os.fdopen(os.dup(fd), "rb") as stream:
                data = stream.read(self.config.max_source + 1)
            if len(data) > self.config.max_source:
                raise _fail()
            return data
        except GatewayError:
            raise
        except Exception as exc:
            raise _fail() from exc
        finally:
            for fd in reversed(descriptors):
                os.close(fd)

    def register(self, path: str) -> str:
        candidate = Path(path)
        if not candidate.is_absolute() or candidate.suffix.lower() not in (".txt", ".pdf", ".docx", ".xlsx"):
            raise _fail()
        for index, root in enumerate(self.config.roots):
            try:
                relative = candidate.relative_to(root)
            except ValueError:
                continue
            self._source_bytes(index, str(relative))
            document_id = "doc_" + secrets.token_urlsafe(18)
            sealed = self._seal(json.dumps([index, str(relative)]).encode(), document_id)
            with self._db() as db:
                db.execute("INSERT INTO documents VALUES (?, ?, 'source', 1)", (document_id, sealed))
            return document_id
        raise _fail()

    def _row(self, document_id: str) -> tuple[bytes, str, int]:
        if not re.fullmatch(r"(?:doc|out)_[A-Za-z0-9_-]{20,}", document_id):
            raise _fail()
        with self._db() as db:
            row = db.execute("SELECT source,kind,version FROM documents WHERE id=?", (document_id,)).fetchone()
        if row is None:
            raise _fail()
        return row

    def _sanitize(self, text: str) -> str:
        if len(text) > self.config.max_output:
            raise _fail()
        try:
            spans = self.detector.detect(text, self.config.language, self.config.entities, self.config.threshold)
        except Exception as exc:
            raise GatewayError("PII detection unavailable") from exc
        spans = sorted(spans, key=lambda x: (x[0], -(x[1] - x[0])))
        merged = []
        for start, stop, entity in spans:
            if start < 0 or stop > len(text) or stop <= start or entity not in self.config.entities:
                raise GatewayError("PII detection unavailable")
            if merged and start < merged[-1][1]:
                previous_start, previous_stop, _ = merged[-1]
                merged[-1] = (previous_start, max(stop, previous_stop), "PII")
            else:
                merged.append((start, stop, entity))
        result = []
        end = 0
        for start, stop, entity in merged:
            value = text[start:stop]
            digest = hmac.new(self.key, (self.config.scope + "\0" + entity + "\0" + value).encode(), hashlib.sha256).hexdigest()[:20]
            result.extend((text[end:start], f"<{entity}_{digest}>"))
            end = stop
        result.append(text[end:])
        sanitized = "".join(result)
        if len(sanitized) > self.config.max_output:
            raise _fail()
        return sanitized

    def _content(self, document_id: str) -> tuple[str, int, str]:
        sealed, kind, version = self._row(document_id)
        if kind == "output":
            text = self._open_seal(sealed, document_id + ":" + str(version)).decode("utf-8")
            # Reapply current policy. Existing pseudonyms are opaque text.
            return self._sanitize(text), version, "text/plain"
        index, relative = json.loads(self._open_seal(sealed, document_id))
        raw = self._source_bytes(index, relative)  # authorize/read before cache lookup
        suffix = Path(relative).suffix.lower()
        digest = hashlib.sha256(raw).hexdigest()
        parser_packages = {".txt": (), ".pdf": ("pypdf",), ".docx": ("python-docx",), ".xlsx": ("openpyxl",)}
        parser_versions = []
        for package in parser_packages[suffix]:
            try:
                parser_versions.append(metadata_version(package))
            except PackageNotFoundError:
                parser_versions.append("unavailable")
        cache_material = json.dumps([document_id, digest, self.config.policy_digest(), self.detector.identity,
                                     self.config.scope, suffix, parser_versions, "parser-v1"], sort_keys=True).encode()
        cache_key = hashlib.sha256(cache_material).hexdigest()
        with self._db() as db:
            cached = db.execute("SELECT payload,expires FROM cache WHERE cache_key=?", (cache_key,)).fetchone()
        if cached and cached[1] > int(time.time()):
            return self._open_seal(cached[0], cache_key).decode("utf-8"), version, "text/plain"
        extracted = _extract(raw, suffix) if suffix == ".txt" else _extract_bounded(raw, suffix, self.config.max_output, self.config.max_source)
        if not extracted.strip():
            raise _fail()
        sanitized = self._sanitize(extracted)
        with self._db() as db:
            db.execute("INSERT OR REPLACE INTO cache VALUES (?,?,?)", (cache_key, self._seal(sanitized.encode(), cache_key), int(time.time()) + self.config.ttl))
            db.execute("DELETE FROM cache WHERE expires < ?", (int(time.time()),))
        return sanitized, version, "text/plain"

    def read_file(self, document_id: str) -> dict:
        text, version, mime = self._content(document_id)
        return {"id": document_id, "version": version, "mime_type": mime, "content": text}

    def list_files(self) -> list[dict]:
        with self._db() as db:
            rows = db.execute("SELECT id,kind,version FROM documents ORDER BY id").fetchall()
        return [{"id": i, "kind": k, "version": v} for i, k, v in rows]

    def metadata(self, document_id: str) -> dict:
        _, kind, version = self._row(document_id)
        if kind == "source":
            self._content(document_id)
        return {"id": document_id, "kind": kind, "version": version, "mime_type": "text/plain"}

    def search(self, query: str) -> list[dict]:
        if not query or len(query) > 256:
            raise _fail()
        results = []
        for row in self.list_files():
            try:
                text, _, _ = self._content(row["id"])
            except GatewayError:
                continue
            position = text.casefold().find(query.casefold())
            if position >= 0:
                results.append({"id": row["id"], "excerpt": text[max(0, position - 60):position + len(query) + 60]})
        return results

    def create(self, content: str) -> dict:
        if not isinstance(content, str):
            raise _fail()
        sanitized = self._sanitize(content)
        document_id = "out_" + secrets.token_urlsafe(18)
        with self._db() as db:
            db.execute("INSERT INTO documents VALUES (?, ?, 'output', 1)", (document_id, self._seal(sanitized.encode(), document_id + ":1")))
        return {"id": document_id, "version": 1, "mime_type": "text/plain"}

    def update(self, document_id: str, content: str, expected_version: int) -> dict:
        if not isinstance(content, str):
            raise _fail()
        sanitized = self._sanitize(content)
        with self._db() as db:
            cursor = db.execute("UPDATE documents SET source=?,version=version+1 WHERE id=? AND kind='output' AND version=?",
                                (self._seal(sanitized.encode(), document_id + ":" + str(expected_version + 1)), document_id, expected_version))
            if cursor.rowcount != 1:
                raise _fail()
        return {"id": document_id, "version": expected_version + 1, "mime_type": "text/plain"}

    def export(self, document_id: str) -> dict:
        text, version, _ = self._content(document_id)
        export_dir = self.config.workspace / "exports"
        export_dir.mkdir(mode=0o700, exist_ok=True)
        filename = document_id + "_v" + str(version) + "_" + secrets.token_hex(6) + ".txt"
        directory_fd = os.open(export_dir, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)
        temporary = ".pending-" + secrets.token_hex(16)
        try:
            fd = os.open(temporary, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o600, dir_fd=directory_fd)
            with os.fdopen(fd, "wb") as stream:
                stream.write(text.encode())
                stream.flush()
                os.fsync(stream.fileno())
            os.rename(temporary, filename, src_dir_fd=directory_fd, dst_dir_fd=directory_fd)
        finally:
            try:
                os.unlink(temporary, dir_fd=directory_fd)
            except FileNotFoundError:
                pass
            os.close(directory_fd)
        return {"id": document_id, "version": version, "export_name": filename, "mime_type": "text/plain"}
